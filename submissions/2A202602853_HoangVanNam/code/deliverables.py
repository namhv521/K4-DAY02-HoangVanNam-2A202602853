"""Export the required workbook, measured plots, report and submission bundle."""
from __future__ import annotations
import json
from pathlib import Path
import shutil
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image
from eval import read_pred, compute_metrics
import dataset


def table(rows, columns):
    return pd.DataFrame(rows).reindex(columns=columns)


def text_table(frame):
    if frame.empty:
        return 'Chưa có kết quả đo cho phần này.\n'
    def cell(value):
        if isinstance(value, (float, np.floating)):
            return 'n.a.' if pd.isna(value) else f'{value:.4f}'
        return str(value).replace('|', '/')
    lines = ['| ' + ' | '.join(frame.columns) + ' |', '| ' + ' | '.join(['---'] * len(frame.columns)) + ' |']
    lines += ['| ' + ' | '.join(cell(v) for v in row) + ' |' for row in frame.itertuples(index=False, name=None)]
    return '\n'.join(lines) + '\n'


def final_tables(output, state):
    final_rows, per_class = [], []
    for tag in ('T00', 'F01'):
        predictions = sorted((output / 'predictions').glob(f'{tag}_seed*_test.csv'))
        measured = []
        for path in predictions:
            pred = read_pred(str(path))
            metrics = compute_metrics(pred.y_true, pred.y_pred, pred.probs)
            val_path = output / 'predictions' / f'{tag}_seed{pred.seed}_val.csv'
            val = read_pred(str(val_path))
            vm = compute_metrics(val.y_true, val.y_pred, val.probs)
            measured.append(metrics)
            final_rows.append({'exp_id': tag, 'seed': pred.seed, 'configuration': json.dumps(state.get('sealed', {}), ensure_ascii=False),
                               'macro_f1_val': vm['macro_f1'], 'macro_f1_test': metrics['macro_f1'],
                               'top1_test': metrics['top1'], 'ece_test': metrics['ece']})
        if measured:
            aggregate = {'exp_id': tag, 'seed': 'mean ± std (ddof=1)'}
            for key in ('macro_f1', 'top1', 'ece'):
                values = [m[key] for m in measured]
                sd = np.std(values, ddof=1) if len(values) >= 2 else float('nan')
                aggregate[key + '_test'] = f'{np.mean(values):.4f} ± {sd:.4f}'
            final_rows.append(aggregate)
            for label, name in enumerate(dataset.CLASS_NAMES):
                row = {'exp_id': tag, 'class': name, 'test_images_per_seed': int(measured[0]['support'][label])}
                for key in ('precision', 'recall', 'f1'):
                    values = [m[key][label] for m in measured]
                    row[key] = float(np.mean(values))
                    row[key + '_std'] = float(np.std(values, ddof=1)) if len(values) > 1 else None
                per_class.append(row)
            cm = sum(m['confusion'] for m in measured)
            fig, ax = plt.subplots(figsize=(10, 9))
            plot = ax.imshow(cm, cmap='Blues')
            ax.set(xticks=range(9), yticks=range(9), xticklabels=dataset.CLASS_NAMES,
                   yticklabels=dataset.CLASS_NAMES, xlabel='Predicted', ylabel='True', title=f'{tag}: summed test confusion, {len(measured)} seeds')
            plt.setp(ax.get_xticklabels(), rotation=65, ha='right')
            for i in range(9):
                for j in range(9):
                    ax.text(j, i, str(cm[i, j]), ha='center', va='center', fontsize=7,
                            color='white' if cm[i, j] > cm.max()/2 else 'black')
            fig.colorbar(plot, ax=ax)
            fig.tight_layout()
            fig.savefig(output / 'curves' / f'{tag}_confusion.png', dpi=150)
            plt.close(fig)
    return table(final_rows, ['exp_id', 'configuration', 'seed', 'macro_f1_val', 'macro_f1_test', 'top1_test', 'ece_test']), table(per_class,
        ['exp_id', 'class', 'test_images_per_seed', 'precision', 'precision_std', 'recall', 'recall_std', 'f1', 'f1_std'])


def error_gallery(output, state):
    path = output / 'predictions' / 'F01_seed0_test.csv'
    if not path.exists():
        return []
    pred = read_pred(str(path))
    wrong = np.flatnonzero(pred.y_true != pred.y_pred)
    indices = sorted(wrong, key=lambda i: -pred.probs[i].max())[:9]
    if not indices:
        return []
    fig, axes = plt.subplots(3, 3, figsize=(12, 12))
    cfg = state['sealed']['final_config']
    rows = []
    for ax in axes.flat:
        ax.axis('off')
    for ax, i in zip(axes.flat, indices):
        filename = pred.filenames[i]
        with Image.open(Path(cfg['images_dir']) / filename) as image:
            ax.imshow(image.convert('RGB'))
        true, guess = dataset.CLASS_NAMES[pred.y_true[i]], dataset.CLASS_NAMES[pred.y_pred[i]]
        ax.set_title(f'{filename}\nTrue: {true}\nPred: {guess}; confidence={pred.probs[i].max():.3f}', fontsize=9)
        rows.append({'filename': filename, 'true': true, 'predicted': guess, 'confidence': float(pred.probs[i].max())})
    fig.tight_layout()
    fig.savefig(output / 'curves' / 'F01_errors.png', dpi=150)
    plt.close(fig)
    pd.DataFrame(rows).to_csv(output / 'eval_out' / 'F01_errors.csv', index=False)
    return rows


def export(output, state):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    (output / 'curves').mkdir(exist_ok=True)
    (output / 'eval_out').mkdir(exist_ok=True)
    backbone_rows = []
    for r in state['Backbones']:
        c = r['config']
        backbone_rows.append(dict(exp_id=r['exp_id'], backbone=r['backbone'], weight_tag=r['weight_tag'],
            params_m=r['params_m'], gmacs=r['gmacs'], img_size=c['img_size'], epochs=c['epochs'], seed=r['seed'],
            macro_f1_val=r['macro_f1_val'], top1_val=r.get('top1_val'),
            train_seconds_per_epoch=r['mean_epoch_seconds'], batch1_ms=r['latency_ms'], notes=r['gmacs_note']))
    backbones = table(backbone_rows, ['exp_id', 'backbone', 'weight_tag', 'params_m', 'gmacs', 'img_size', 'epochs', 'seed',
                    'macro_f1_val', 'top1_val', 'train_seconds_per_epoch', 'batch1_ms', 'notes'])
    anchor = next((r['macro_f1_val'] for r in state['Training'] if r['exp_id'] == 'T00'), None)
    training_rows = []
    for r in state['Training']:
        pred = read_pred(str(output / 'predictions' / f"{r['exp_id']}_seed{r['seed']}_val.csv"))
        m = compute_metrics(pred.y_true, pred.y_pred, pred.probs)
        training_rows.append(dict(exp_id=r['exp_id'], backbone=r['backbone'], axis=r['axis'],
            changes=json.dumps(r['changes'], ensure_ascii=False), seed=r['seed'], macro_f1_val=r['macro_f1_val'],
            top1_val=m['top1'], delta_vs_T00=r['macro_f1_val'] - anchor,
            f1_chinee_apple=m['f1'][0], f1_snake_weed=m['f1'][7],
            notes='One screening seed; differences do not establish statistical significance'))
    training = table(training_rows, ['exp_id', 'backbone', 'axis', 'changes', 'seed', 'macro_f1_val', 'top1_val',
                                   'delta_vs_T00', 'f1_chinee_apple', 'f1_snake_weed', 'notes'])
    inference_rows = [dict(r) for r in state['Inference']]
    reference = next((r['p50'] for r in inference_rows if r['method'] == 'single'), None)
    for r in inference_rows:
        r['relative_cost'] = r['p50'] / reference
    inference = table(inference_rows, ['exp_id', 'method', 'checkpoint', 'K', 'macro_f1_val', 'top1_val', 'ece_val',
                      'temperature', 'ece_calibrated_val', 'p50', 'p95', 'p99', 'bulk_images_per_s', 'relative_cost'])
    final, per_class = final_tables(output, state)
    latency = table(state['Latency'], ['exp_id', 'method', 'checkpoint', 'gpu', 'dtype', 'batch', 'img_size', 'bn_fused',
                                      'p50', 'p95', 'p99', 'images_per_s', 'preprocessing', 'torch'])
    ranking = [{'exp_id': r['exp_id'], 'stage': stage, 'backbone': r['backbone'], 'macro_f1_val': r['macro_f1_val'],
                'batch1_p50_ms': r.get('latency_ms'), 'seed': r['seed']} for stage in ('Backbones', 'Training') for r in state[stage]]
    ranking += [{'exp_id': r['exp_id'], 'stage': 'Inference', 'backbone': r['config']['backbone'],
                 'macro_f1_val': r['macro_f1_val'], 'batch1_p50_ms': r['p50'], 'seed': 0} for r in state['Inference']]
    summary = table(sorted(ranking, key=lambda r: -r['macro_f1_val'])[:10],
                    ['exp_id', 'stage', 'backbone', 'macro_f1_val', 'batch1_p50_ms', 'seed'])
    sheets = {'Summary': summary, 'Backbones': backbones, 'Training': training, 'Inference': inference,
              'Final': final, 'PerClass': per_class, 'Latency': latency}
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(output / 'results.xlsx', engine='openpyxl') as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False, startrow=3)
            ws = writer.sheets[name]
            ws['A1'] = f"{state['student_id']} — {state['student_name']} — {name}"
            ws['A1'].font = Font(bold=True, size=14)
            ws['A2'] = 'Source: experiment_state.json, runs/*/history.csv and original eval.py; unavailable values remain blank.'
            ws.freeze_panes = 'C5'
            ws.auto_filter.ref = f'A4:{get_column_letter(len(frame.columns))}{max(4, len(frame)+4)}'
            ws.sheet_view.showGridLines = False
            for cell in ws[4]:
                cell.font = Font(color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='17365D')
                cell.alignment = Alignment(wrap_text=True, vertical='center')
            ws.row_dimensions[4].height = 42
            for col, label in enumerate(frame.columns, 1):
                ws.column_dimensions[get_column_letter(col)].width = min(36, max(15, len(label)+3))
            for row in ws.iter_rows(min_row=5):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
                    if isinstance(cell.value, float):
                        cell.number_format = '0.0000'
            if frame.empty:
                ws['A5'] = 'NOT RUN — no measured results for this stage yet.'
                ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=len(frame.columns))
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(frame.columns))
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(frame.columns))
            ws['A2'].alignment = Alignment(wrap_text=True)
            ws.row_dimensions[2].height = 32
            if 'macro_f1_val' in frame and not frame.empty:
                row_number = int(frame['macro_f1_val'].astype(float).idxmax()) + 5
                for cell in ws[row_number]:
                    cell.fill = PatternFill('solid', fgColor='E2F0D9')
    if not inference.empty:
        fig, ax = plt.subplots(figsize=(9, 5))
        ax.scatter(inference.p95, inference.macro_f1_val)
        for r in inference.itertuples():
            ax.annotate(r.method, (r.p95, r.macro_f1_val), fontsize=8)
        ax.set(xlabel='p95 latency (ms, batch 1)', ylabel='Validation macro-F1', title='Accuracy / latency trade-off')
        fig.tight_layout()
        fig.savefig(output / 'curves' / 'inference_tradeoff.png', dpi=150)
        plt.close(fig)
    if not backbones.empty:
        fig, ax = plt.subplots(figsize=(9, 5))
        ax.scatter(backbones.batch1_ms, backbones.macro_f1_val)
        for r in backbones.itertuples():
            ax.annotate(r.backbone, (r.batch1_ms, r.macro_f1_val), fontsize=8)
        ax.set(xlabel='Latency (ms, batch 1)', ylabel='Validation macro-F1', title='Backbone screening')
        fig.tight_layout()
        fig.savefig(output / 'curves' / 'backbone_tradeoff.png', dpi=150)
        plt.close(fig)
    errors = error_gallery(output, state)
    root = Path(__file__).resolve().parents[1]
    preparation = root / 'artifacts' / 'preparation'
    if preparation.exists():
        shutil.copytree(preparation, output / 'artifacts' / 'preparation', dirs_exist_ok=True)
    if (root / 'report.md').exists():
        prep_report = (root / 'report.md').read_text(encoding='utf-8')
    else:
        prep_report = (preparation / 'eda.md').read_text(encoding='utf-8') if (preparation / 'eda.md').exists() else 'Run preparation first.'
        prep_report = prep_report.replace('(class_', '(artifacts/preparation/class_')
    lines = [f"# Báo cáo Lab Day 2 — {state['student_name']}", f"Mã học viên: **{state['student_id']}**", '',
        '## 1. Tóm tắt',
        f"Đã ghi nhận {len(backbones)} backbone, {len(training)} công thức huấn luyện, {len(inference)} phương pháp suy luận và {len(state['Final'])} seed chung kết.",
        'Các bảng dưới đây chỉ lấy từ dữ liệu đã đo; phần chưa chạy không có số liệu.', '',
        '## 2. Dữ liệu và thiết lập', prep_report, '',
        'Các thí nghiệm chính dùng fold 0 cố định, chọn checkpoint bằng macro-F1 validation. '
        'T00 là công thức nền trên backbone được chọn từ Bước 1; B01_resnet50 giữ mốc kiến trúc ResNet-50. '
        'Batch mặc định 32 để phù hợp GPU T4 (giảm từ 64 gợi ý trong GUIDE); 12 epoch; AdamW, warmup 1 epoch, cosine; AMP train trên CUDA. '
        'Nhãn train không sửa theo labels.csv khi có sai lệch upstream đã ghi trong EDA.', '',
        '## 3. So sánh backbone', text_table(backbones[['backbone', 'macro_f1_val', 'top1_val', 'params_m', 'gmacs', 'batch1_ms']]),
        'Backbone có macro-F1 val cao nhất được chọn. Đây là sàng lọc một seed; trọng số tiền huấn luyện '
        'có thể có công thức khác nhau nên không quy mọi chênh lệch cho riêng kiến trúc. GMAC bị để trống nếu fvcore gặp phép toán chưa hỗ trợ.', '',
        '## 4. Ablation công thức', text_table(training[['exp_id', 'axis', 'changes', 'macro_f1_val', 'delta_vs_T00']]),
        'Mỗi dòng A–G thay một yếu tố; T99 kết hợp các thay đổi có delta dương trên từng trục. '
        'Kết hợp không mặc nhiên tốt hơn: vẫn được đo lại và so với tất cả công thức. '
        'Không kết luận ý nghĩa thống kê từ một seed sàng lọc. Label smoothing giảm độ tự tin; focal ưu tiên mẫu khó; '
        'loss có trọng số và sampler tác động khác nhau lên lớp hiếm. Đây là cơ chế lý thuyết, hiệu quả thực tế phải đọc từ bảng.', '',
        '## 5. Phương pháp suy luận', text_table(inference[['method', 'macro_f1_val', 'ece_val', 'ece_calibrated_val', 'p95', 'relative_cost']]),
        '![Đánh đổi](curves/inference_tradeoff.png)' if not inference.empty else 'Chưa có biểu đồ suy luận.',
        'Latency đo 10 warmup + 50 lượt, đồng bộ CUDA trước/sau, batch 1 và 32. Không tính decode ảnh/tiền xử lý PIL hoặc I/O. '
        'Temperature khớp bằng NLL trên validation, bảo toàn argmax; giảm NLL không bảo đảm ECE cũng giảm. '
        'ViT/Swin có kích thước cố định không chạy resize tùy ý. BN fusion chỉ áp dụng kiến trúc có BatchNorm và kiểm tra tương đương.', '',
        '## 6. Chung kết và test', text_table(final),
        'Ba seed 0/1/2 cho cả F01 và T00; mean ± sample std (ddof=1). Mọi lựa chọn được khóa trước lần test đầu tiên. '
        'Test cache cho phép xuất lại báo cáo mà không chạy lại model trên test. Không dùng test để chọn hoặc sửa công thức.', '',
        '## 7. Phân tích lỗi và khuyến nghị']
    if state['Final']:
        fv = np.array([r['macro_f1_test'] for r in state['Final']])
        bv = np.array([r['baseline']['macro_f1'] for r in state['Final']])
        noise = max(np.std(fv, ddof=1), np.std(bv, ddof=1))
        lines += [f'Macro-F1: F01 {fv.mean():.4f} ± {fv.std(ddof=1):.4f}; T00 {bv.mean():.4f} ± {bv.std(ddof=1):.4f}; delta {fv.mean()-bv.mean():+.4f}.',
                  'Chênh lệch lớn hơn std đo được; vẫn chưa phải kiểm định ý nghĩa thống kê.' if abs(fv.mean()-bv.mean()) > noise else 'Chênh lệch không vượt std: chưa phân biệt được hai cấu hình.',
                  '![Ma trận nhầm lẫn](curves/F01_confusion.png)']
        if errors:
            lines += ['![Ảnh dự đoán sai, seed 0](curves/F01_errors.png)', text_table(pd.DataFrame(errors)),
                      'Các cặp nhầm lẫn và ảnh trên là quan sát thực tế. Nền lá/cỏ tương tự, che khuất và ánh sáng là giả thuyết cần kiểm tra bằng mắt; không tự động khẳng định nguyên nhân sinh học.']
    else:
        lines += ['Chưa chạy chung kết; chưa có kết luận về chất lượng test.']
    if not inference.empty:
        for budget in (30, 100):
            eligible = inference[inference.p95 <= budget]
            lines += [f'Ngân sách {budget} ms: ' + (f"chọn {eligible.sort_values('macro_f1_val', ascending=False).iloc[0]['method']} theo macro-F1 validation trong các cấu hình đáp ứng p95." if not eligible.empty else 'không có cấu hình đo được đáp ứng trên phần cứng này.')]
    lines += ['', '## 8. Hạn chế',
              'Một fold, ba seed vòng cuối và một seed sàng lọc. Split ngẫu nhiên theo ảnh có thể lạc quan khi chuyển địa điểm/mùa. '
              'Chưa đo end-to-end trên robot; độ trễ Colab không thay thế latency thiết bị triển khai. '
              'Calibration dùng cùng val chọn công thức, có thể lạc quan; miền mới cần đánh giá riêng. '
              'Không đối chiếu trực tiếp số của bài báo 100 epoch với thiết lập 12 epoch này.', '',
              '## 9. Phụ lục tái lập',
              'Cấu hình, tag trọng số, phiên bản thư viện và phần cứng: `runs/*/seed*/config.json`. '
              'Danh sách công thức và lựa chọn khóa: `experiment_state.json`. Notebook tự chứa toàn bộ code: `code/lab_day2.ipynb`. '
              'Kết quả chấm gốc: `eval_out/`; các file dự đoán có p0…p8, argmax và nhãn thật.']
    (output / 'report.md').write_text('\n\n'.join(lines) + '\n', encoding='utf-8')
    source = Path(__file__).resolve().parent
    target = output / 'code'
    if source.resolve() != target.resolve():
        shutil.copytree(source, target, dirs_exist_ok=True, ignore=shutil.ignore_patterns('__pycache__', '*.pyc'))
    shutil.copy2(root / 'eval.py', output / 'eval.py')
    for name in ('dochieu.md',):
        if (root / name).exists():
            shutil.copy2(root / name, output / name)
    readme = f'''# Lab Day 2 — {state['student_name']} ({state['student_id']})

Mở `code/lab_day2.ipynb` trên Google Colab, chọn GPU rồi Run all.
Notebook tự chứa code, tải DeepWeeds kiểm MD5, lưu Drive và resume theo epoch.
Nếu dừng phiên, mở lại notebook, giữ nguyên cấu hình và chạy lại; không đổi thư mục output sau khi đã mở test.

Các bước: chuẩn bị → backbone → training → inference → final/test → export.
Seed sàng lọc 0; vòng cuối 0,1,2. T00 dùng công thức nền trên backbone được chọn; B01 giữ mốc ResNet-50.
Cấu hình chính: {state.get('run_settings', {})}. Chọn hoàn toàn bằng validation.

`results.xlsx` có 7 sheet yêu cầu. `report.md` được sinh từ số liệu đo.
`runs/` chứa log/cấu hình/checkpoint; không đưa checkpoint hoặc dataset vào ZIP bài nộp.
`eval_out/` chứa kết quả `eval.py score/grade`. Trọng số/thư viện/phần cứng thực tế có trong từng config.json.
`dochieu.md` giải thích từng file.

Notebook chưa có URL chia sẻ công khai: sau khi lưu vào Colab, dùng Share để lấy link của bạn.
'''
    (output / 'README.md').write_text(readme, encoding='utf-8')
    print(f'Exported {output / "results.xlsx"} and report.md from measured data.', flush=True)
